import asyncio
import re
import zlib

import pytest
from docx import Document

from app import config, conversations, documents, gemini, ingest, pipeline, vector_store
from app.schemas import BatchNotes, FinalAnswer
from app.text_utils import normalize_ar, split_message

DIM = 768


def fake_vec(text: str) -> list[float]:
    v = [0.0] * DIM
    for w in re.findall(r"\w+", normalize_ar(text)):
        v[zlib.crc32(w.encode()) % DIM] += 1.0
    n = sum(x * x for x in v) ** 0.5 or 1.0
    return [x / n for x in v]


@pytest.fixture(autouse=True)
def env(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "QDRANT_URL", ":memory:")
    monkeypatch.setattr(config, "COLLECTION_NAME", f"t_{tmp_path.name[-8:]}")
    monkeypatch.setattr(config, "EMBED_DIM", DIM)
    monkeypatch.setattr(config, "CONVERSATIONS_DIR", str(tmp_path / "conversations"))
    monkeypatch.setattr(config, "PARSED_DIR", str(tmp_path / "parsed"))
    monkeypatch.setattr(config, "CHUNK_MAX_CHARS", 320)
    monkeypatch.setattr(config, "CHUNK_MIN_CHARS", 60)
    monkeypatch.setattr(vector_store, "_client", None)

    async def fake_embed(texts, task_type):
        assert task_type in ("RETRIEVAL_DOCUMENT", "RETRIEVAL_QUERY")
        return [fake_vec(t) for t in texts]

    monkeypatch.setattr(gemini, "embed", fake_embed)


def make_docx(path):
    d = Document()
    d.add_heading("لائحة كلية الذكاء الاصطناعي", level=1)
    d.add_paragraph("مادة (3) الغياب: يحرم الطالب من دخول الامتحان إذا تجاوزت نسبة غيابه 25% من المحاضرات.")
    d.add_paragraph("مادة (4) الرسوم: تسدد الرسوم الدراسية قبل بداية الفصل الدراسي بأسبوعين.")
    t = d.add_table(rows=3, cols=3)
    for r, row in enumerate([("المادة", "الدكتور", "اليوم"), ("تعلم الآلة", "د. أحمد سمير", "السبت"), ("الرؤية الحاسوبية", "د. منى علي", "الاثنين")]):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    d.save(path)


# ============== نصوص وتقطيع ==============
def test_normalize_and_split():
    assert normalize_ar("الإِمْتِحَانَات") == normalize_ar("الامتحانات")
    assert normalize_ar("مدرسة") == normalize_ar("مدرسه")
    parts = split_message("سطر\n" * 3000, limit=4000)
    assert len(parts) > 1 and all(len(p) <= 4000 for p in parts)


def test_page_markers_roundtrip():
    pages = [(3, "نص الصفحة ثلاثة"), (4, "نص الصفحة أربعة")]
    assert documents.parse_page_markers(documents.pages_to_markdown(pages)) == pages


def test_chunking_keeps_articles_and_pages():
    pages = [(7, "# مادة (5) شروط القيد\n\n" + "كلمة " * 120), (8, "مادة (6) حذف واضافة\n\nنص قصير عن الحذف.")]
    chunks = documents.chunk_pages(pages, "reg.pdf")
    assert len(chunks) >= 2
    assert all(len(c.text) <= config.CHUNK_MAX_CHARS + 50 for c in chunks)
    assert chunks[0].page_start == 7 and "مادة (5)" in chunks[0].section
    assert any("مادة (6)" in c.text and c.page_start in (7, 8) for c in chunks)


def test_long_table_split_repeats_header():
    rows = "\n".join(f"| مادة {i} | د. فلان {i} | السبت |" for i in range(60))
    chunks = documents.chunk_pages([(1, "| المادة | الدكتور | اليوم |\n| --- | --- | --- |\n" + rows)], "t.docx")
    assert len(chunks) > 1
    assert all(c.text.startswith("| المادة | الدكتور | اليوم |") for c in chunks)


# ============== الرفع (DOCX حقيقي) ==============
def test_ingest_docx_and_idempotent(tmp_path):
    async def go():
        p = tmp_path / "reg.docx"
        make_docx(str(p))
        await vector_store.init_collection()
        n1 = await ingest.ingest_file(str(p), "reg.docx", force=True)
        n2 = await ingest.ingest_file(str(p), "reg.docx", force=True)
        sources = await vector_store.list_sources()
        assert n1 == n2 and sources["reg.docx"] == n1
        md = (tmp_path / "parsed" / "reg.docx.md").read_text(encoding="utf-8")
        assert "أحمد سمير" in md and "| --- |" in md
        await vector_store.delete_source("reg.docx")
        assert not await vector_store.list_sources()

    asyncio.run(go())


# ============== الخط كامل ==============
def install_fake_llm(monkeypatch, calls, answer="الرد النهائي"):
    async def fake_json(prompt, schema, model=None):
        calls["json"].append(schema)
        calls["prompt"] = prompt
        if schema is FinalAnswer:
            return {"answer": answer, "location_ids": []}
        assert schema is BatchNotes
        notes = []
        if "25%" in prompt:
            notes.append({"fact": "الغياب 25%", "evidence": "نسبة غيابه 25%", "source": "reg.docx", "location_ids": []})
        return {"notes": notes}

    monkeypatch.setattr(gemini, "generate_json", fake_json)


async def setup_docx(tmp_path):
    p = tmp_path / "reg.docx"
    make_docx(str(p))
    await vector_store.init_collection()
    await ingest.ingest_file(str(p), "reg.docx", force=True)


def test_single_llm_call_with_chunks_and_history(monkeypatch, tmp_path):
    calls = {"json": []}
    install_fake_llm(monkeypatch, calls)

    async def go():
        await setup_docx(tmp_path)
        answer = await pipeline.answer_message(1, "كام نسبة الغياب؟ ومين بيدرس تعلم الآلة؟")
        assert answer == "الرد النهائي"
        assert calls["json"] == [FinalAnswer]
        assert "25%" in calls["prompt"] and "أحمد سمير" in calls["prompt"]
        hist = await conversations.get_history(1)
        assert [m["role"] for m in hist] == ["user", "bot"]

    asyncio.run(go())


def test_conversation_and_follow_up_query(monkeypatch, tmp_path):
    calls = {"json": []}
    install_fake_llm(monkeypatch, calls)
    queries = []
    real_embed = gemini.embed

    async def spy_embed(texts, task_type):
        if task_type == "RETRIEVAL_QUERY":
            queries.append(list(texts))
        return await real_embed(texts, task_type)

    monkeypatch.setattr(gemini, "embed", spy_embed)

    async def go():
        await setup_docx(tmp_path)
        await pipeline.answer_message(7, "عايز أعرف الغياب")
        await pipeline.answer_message(7, "طب والرسوم؟")
        assert "الطالب: عايز أعرف الغياب" in calls["prompt"] and "البوت: الرد النهائي" in calls["prompt"]
        assert queries[0] == ["عايز أعرف الغياب"]
        assert queries[1] == ["طب والرسوم؟", "عايز أعرف الغياب طب والرسوم؟"]

    asyncio.run(go())


def test_large_context_goes_through_batches(monkeypatch, tmp_path):
    calls = {"json": []}
    install_fake_llm(monkeypatch, calls)
    monkeypatch.setattr(config, "BATCH_CHAR_BUDGET", 150)
    monkeypatch.setattr(config, "BATCH_MAX_CHUNKS", 1)

    async def go():
        await setup_docx(tmp_path)
        assert await pipeline.answer_message(3, "كام نسبة الغياب؟") == "الرد النهائي"
        assert calls["json"].count(BatchNotes) >= 2 and calls["json"][-1] is FinalAnswer
        assert "الغياب 25%" in calls["prompt"]

    asyncio.run(go())


def test_answer_failure_returns_busy_and_is_not_saved(monkeypatch, tmp_path):
    async def boom(*a, **k):
        raise RuntimeError("quota")

    monkeypatch.setattr(gemini, "generate_json", boom)

    async def go():
        await setup_docx(tmp_path)
        assert await pipeline.answer_message(2, "ما الغياب؟") == pipeline.MSG_BUSY
        assert [m["role"] for m in await conversations.get_history(2)] == ["user"]

    asyncio.run(go())


def test_all_batches_failing_returns_busy(monkeypatch, tmp_path):
    async def fake_json(prompt, schema, model=None):
        if schema is BatchNotes:
            raise RuntimeError("quota")
        return {"answer": "x", "location_ids": []}

    monkeypatch.setattr(gemini, "generate_json", fake_json)
    monkeypatch.setattr(config, "BATCH_CHAR_BUDGET", 150)
    monkeypatch.setattr(config, "BATCH_MAX_CHUNKS", 1)

    async def go():
        await setup_docx(tmp_path)
        assert await pipeline.answer_message(4, "ما الغياب؟") == pipeline.MSG_BUSY

    asyncio.run(go())


def test_empty_answer_returns_not_found(monkeypatch, tmp_path):
    install_fake_llm(monkeypatch, {"json": []}, answer="   ")

    async def go():
        await setup_docx(tmp_path)
        assert await pipeline.answer_message(5, "سؤال") == pipeline.MSG_NOT_FOUND

    asyncio.run(go())


def test_interleave_gives_each_query_a_share():
    mk = lambda i: pipeline.Hit(str(i), "t", "s", 1, 1, "", 1.0)
    out = pipeline.interleave([[mk(1), mk(2), mk(3)], [mk(10), mk(11)]], limit=4)
    assert [h.id for h in out] == ["1", "10", "2", "11"]


def test_lexical_terms_drop_stopwords():
    terms = pipeline.lexical_terms("فين كلية الطب البيطري؟")
    assert normalize_ar("كلية") in terms and normalize_ar("فين") not in terms
